"""
大纲→Word 生成模块：内容生成与格式渲染分离

职责说明：
- 内容生成：LLM 根据 title + outline 逐节生成内容，
  每节输出结构化文本（标题 + 若干段落 + 列表）
- 大纲层级：条目支持 markdown 前缀（#/##/### 对应 1/2/3 级，无前缀视为 1 级）；
  层级越深篇幅越短（1 级 3-5 段，2/3 级 1-2 段），控制总量与 token
- 写作要求注入：条目下一行以 > 开头视为该节的写作要求
  （如「> 要体现互相认识的目的」），注入该节 Prompt 的「特别要求」
- 模板渲染：python-docx 渲染，模板样式：
  - 文档标题：黑体 16pt 加粗
  - 一级标题：黑体 14pt 加粗
  - 二级标题：黑体 13pt 加粗
  - 三级标题：黑体 12pt 加粗
  - 正文：宋体 12pt，1.5 倍行距
  - 列表：项目符号列表
- 生成后写入 PROCESSED_DIR
- 流式：generate_document_stream 逐节产出事件（section_start / section_done / done / error），
  供 /ai/generate SSE 接口使用（长文档不再受接口总超时限制）
"""
import os
import re
import uuid
from dataclasses import dataclass, field

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from app.llm.client import LLMClient

# 大纲条目层级解析：'# 章' / '## 节' / '### 小节'；无前缀视为 1 级
_OUTLINE_RE = re.compile(r"^(#{1,3})\s*(.+)$")


def parse_outline_item(item: str) -> tuple[int, str]:
    """解析一条大纲条目 → (层级, 标题)。空标题返回 (层级, '')，由调用方跳过"""
    m = _OUTLINE_RE.match(item.strip())
    if m:
        return len(m.group(1)), m.group(2).strip()
    return 1, item.strip()


def parse_outline_with_requirements(outline: list[str]) -> list[tuple[int, str, str]]:
    """解析大纲，返回 [(层级, 标题, 写作要求)]。

    - 条目行：与 parse_outline_item 同规则
    - 要求行：紧随条目之后、以 > 开头，注入该节的「特别要求」
      （如「> 要体现互相认识的目的」）；多条要求用「；」合并
    - 条目之前孤立的要求行忽略；空行忽略
    """
    items: list[tuple[int, str, str]] = []
    for raw in outline:
        line = raw.strip()
        if not line:
            continue
        if line.startswith(">"):
            req = line.lstrip(">").strip()
            if items and req:
                level, title, old_req = items[-1]
                items[-1] = (level, title, f"{old_req}；{req}" if old_req else req)
            continue
        level, title = parse_outline_item(line)
        if title:
            items.append((level, title, ""))
    return items


@dataclass
class SectionContent:
    """一节文档内容：层级 + 标题 + 若干段落 + 要点列表"""

    title: str
    level: int = 1
    paragraphs: list[str] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)


class DocxRenderer:
    """模板渲染器：把结构化内容渲染成 Word 文档（python-docx）"""

    # 各级标题字号（pt）；文档大标题固定 16pt
    HEADING_SIZES = {1: 14, 2: 13, 3: 12}

    @staticmethod
    def _style_run(run, font_name: str, size: int, bold: bool = False) -> None:
        """设置字体样式。中文字体必须同时设置 ascii 与 eastAsia，否则中文不生效"""
        run.font.name = font_name
        run.font.size = Pt(size)
        run.font.bold = bold
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

    def render(self, title: str, sections: list[SectionContent]) -> Document:
        """渲染完整文档：文档大标题（黑体 16pt）+ 各节内容（按层级分级标题）"""
        document = Document()
        self._add_heading(document, title, size=16)  # 文档标题：黑体 16pt
        for section in sections:
            self._add_heading(document, section.title, size=self.HEADING_SIZES[section.level])
            for para in section.paragraphs:
                self._add_paragraph(document, para)
            for bullet in section.bullets:
                self._add_bullet(document, bullet)
        return document

    def _add_heading(self, document: Document, text: str, size: int) -> None:
        """添加标题（黑体加粗）"""
        run = document.add_paragraph().add_run(text)
        self._style_run(run, "黑体", size, bold=True)

    def _add_paragraph(self, document: Document, text: str) -> None:
        """添加正文：宋体 12pt，1.5 倍行距"""
        p = document.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        self._style_run(p.add_run(text), "宋体", 12)

    def _add_bullet(self, document: Document, text: str) -> None:
        """添加项目符号列表"""
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.5
        self._style_run(p.add_run(text), "宋体", 12)


class ContentGenerator:
    """内容生成器：LLM 根据 title + outline 逐节生成结构化内容（支持层级大纲）"""

    def __init__(self) -> None:
        self._llm = LLMClient()

    def generate(self, doc_title: str, outline: list[str]) -> list[SectionContent]:
        """对大纲每个条目，调用 LLM 生成一节内容（层级决定篇幅；要求行注入该节 Prompt）"""
        sections: list[SectionContent] = []
        for level, section_title, requirement in parse_outline_with_requirements(outline):
            raw = self._ask_llm(doc_title, section_title, level, requirement)
            sections.append(self._parse_section(section_title, raw, level))
        return sections

    def _ask_llm(
        self,
        doc_title: str,
        section_title: str,
        level: int,
        requirement: str = "",
        context: str = "",
    ) -> str:
        """构造 Prompt 并调用 LLM，返回该节原始文本。

        - level 决定篇幅要求
        - requirement 为写作要求（大纲 > 行 / 重生成弹窗修改意见）
        - context 为前文要点（增量生成时已保留节的概括，保持口径一致）
        """
        if level == 1:
            depth_req = "2. 写 3-5 段正文，每段 2-4 句话；\n3. 正文结束后另起一行，写 2-3 条要点。"
            unit = "章节"
        else:
            depth_req = "2. 写 1-2 段正文，每段 2-3 句话；\n3. 正文结束后另起一行，写 0-2 条要点。"
            unit = "小节"
        req_line = f"4. 特别要求：{requirement}\n" if requirement else ""
        ctx_line = f"5. 前文要点（保持口径一致、避免重复论述）：\n{context}\n" if context else ""
        prompt = (
            f"你是一位文档写作助手。请为文档《{doc_title}》中的{unit}「{section_title}」撰写内容。\n"
            "要求：\n"
            "1. 直接输出正文内容，禁止输出任何思考过程、解释说明或客套话"
            "（如“好的”“以下是”“我将为您”等开场白）；\n"
            f"{depth_req}\n"
            f"{req_line}"
            f"{ctx_line}"
            "格式：要点以 - 开头，其余行均为正文段落。"
        )
        return self._llm.chat([{"role": "user", "content": prompt}])

    @staticmethod
    def _parse_section(title: str, raw: str, level: int) -> SectionContent:
        """解析 LLM 输出：- 开头的行为要点，其余行作为正文段落"""
        paragraphs: list[str] = []
        bullets: list[str] = []
        for line in raw.splitlines():
            line = line.strip()
            if not line:
                continue
            # 过滤掉 LLM 可能输出的元话语/思考痕迹（双保险）
            if _is_meta_line(line):
                continue
            if line.startswith("-"):
                bullets.append(line.lstrip("-").strip())
            else:
                paragraphs.append(line)
        return SectionContent(title=title, level=level, paragraphs=paragraphs, bullets=bullets)


def _is_meta_line(line: str) -> bool:
    """判断一行是否为 LLM 的元话语/思考痕迹（如开场白、解释说明）"""
    meta_prefixes = ("好的", "以下是", "我将", "我来", "下面", "这里", "（思考", "【思考")
    return line.startswith(meta_prefixes) and len(line) < 40


def render_document(title: str, sections: list[SectionContent], dst_dir: str) -> dict:
    """渲染 Word 并落盘 → { filename, filePath }（filename 的 uuid 同时作为对外 fileId）"""
    if not sections:
        raise ValueError("大纲中没有有效条目，无法生成")

    document = DocxRenderer().render(title, sections)

    # uuid 命名避免重名
    filename = f"{uuid.uuid4().hex}.docx"
    dst_path = os.path.join(dst_dir, filename)
    document.save(dst_path)
    return {"filename": filename, "filePath": dst_path}


def generate_document(title: str, outline: list[str], dst_dir: str) -> dict:
    """同步主入口：LLM 生成全部节 → 渲染 Word → 落盘（保留给非流式调用/测试）

    返回：{ filename, filePath }
    """
    if not title or not outline:
        raise ValueError("标题和大纲均不能为空")

    sections = ContentGenerator().generate(title, outline)
    return render_document(title, sections, dst_dir)


def generate_document_stream(title: str, outline: list[str], dst_dir: str, context: str = ""):
    """流式主入口（供 /ai/generate SSE 调用）：逐节 LLM 生成，产出事件 dict 序列。

    事件协议：
        {"type": "section_start", "index", "title"}          每节开始
        {"type": "section_done",  "index", "title", "level", "paragraphs", "bullets"}  每节完成（含全文）
        {"type": "done", "fileId", "filename", "filePath"}   全部完成（已渲染落盘）
        {"type": "error", "message"}                         某节失败（提前结束，不落盘）

    - outline 通常只含「需要新生成的条目」（增量生成由前端 diff 后传入），
      已保留的节不在其中，零重复调用
    - context 为已保留节的前文概括（前端拼接），注入每条新节 Prompt 保持口径一致
    - 某节生成失败：产出 error 事件后停止，已生成的前几节不落盘
      （前端保留已收到的节、对失败节走 /ai/generate/section 单独重试）
    - 每节一个 LLM 调用，调用层自带 120s 超时，流整体不限时
    """
    if not title or not outline:
        raise ValueError("标题和大纲均不能为空")

    items = parse_outline_with_requirements(outline)
    if not items:
        raise ValueError("大纲中没有有效条目，无法生成")

    content = ContentGenerator()
    sections: list[SectionContent] = []
    for idx, (level, section_title, requirement) in enumerate(items):
        yield {"type": "section_start", "index": idx, "title": section_title}
        try:
            raw = content._ask_llm(title, section_title, level, requirement, context)
            section = content._parse_section(section_title, raw, level)
        except Exception as e:
            yield {
                "type": "error",
                "message": f"第 {idx + 1} 节「{section_title}」生成失败: {e}",
            }
            return
        sections.append(section)
        yield {
            "type": "section_done",
            "index": idx,
            "title": section_title,
            "level": section.level,
            "paragraphs": section.paragraphs,
            "bullets": section.bullets,
        }

    result = render_document(title, sections, dst_dir)
    yield {
        "type": "done",
        "fileId": os.path.splitext(result["filename"])[0],
        "filename": result["filename"],
        "filePath": result["filePath"],
    }


def generate_section(
    doc_title: str,
    section_title: str,
    level: int = 1,
    requirement: str = "",
    context: str = "",
) -> dict:
    """单节生成（供 /ai/generate/section 调用）：一次 LLM 调用，返回结构化内容。

    用于：单节重生成（含修改意见）、失败节重试。
    """
    if not doc_title.strip() or not section_title.strip():
        raise ValueError("文档标题与节标题均不能为空")
    level = max(1, min(3, int(level or 1)))
    content = ContentGenerator()
    raw = content._ask_llm(doc_title, section_title, level, requirement, context)
    section = content._parse_section(section_title, raw, level)
    return {
        "title": section.title,
        "level": section.level,
        "paragraphs": section.paragraphs,
        "bullets": section.bullets,
    }


def _coerce_sections(raw: list[dict]) -> list[SectionContent]:
    """把前端传来的分节内容列表校验并转为 SectionContent（防御脏数据：缺字段/空节跳过）"""
    sections: list[SectionContent] = []
    for item in raw or []:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or "").strip()
        if not title:
            continue
        level = max(1, min(3, int(item.get("level") or 1)))
        paragraphs = [str(p).strip() for p in (item.get("paragraphs") or []) if str(p).strip()]
        bullets = [str(b).strip() for b in (item.get("bullets") or []) if str(b).strip()]
        if not paragraphs and not bullets:
            continue  # 空节跳过
        sections.append(
            SectionContent(title=title, level=level, paragraphs=paragraphs, bullets=bullets)
        )
    return sections


def assemble_document(title: str, sections: list[dict], dst_dir: str) -> dict:
    """渲染落盘（不调 LLM）：把分节内容组装成 Word（供 /ai/generate/assemble 调用）。

    前端编辑/重生成/删除后点下载时调用，秒级完成——「内容生成与格式渲染分离」的红利。
    """
    if not title or not title.strip():
        raise ValueError("标题不能为空")
    coerced = _coerce_sections(sections)
    return render_document(title, coerced, dst_dir)
